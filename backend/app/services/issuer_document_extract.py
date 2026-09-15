"""本地上传文件抽取纯文本：PDF / 纯文本 / CSV / XLSX。"""
from __future__ import annotations

import csv
import io
from pathlib import Path

_ROW_CAP = 50_000
_CHAR_CAP = 500_000


def _read_pdf(path: Path, lim: int = 200_000) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        r = PdfReader(str(path))
        out, n = [], 0
        for pg in r.pages[:120]:
            t = pg.extract_text() or ""
            out.append(t)
            n += len(t)
            if n >= lim:
                break
        return "\n".join(out)[:lim]
    except Exception:
        return ""


def _read_plain(path: Path, lim: int = _CHAR_CAP) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")[:lim]


def _read_csv(path: Path, lim: int = _CHAR_CAP) -> str:
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="ignore")
    buf = io.StringIO(text)
    reader = csv.reader(buf)
    lines: list[str] = []
    total = 0
    for i, row in enumerate(reader):
        if i >= _ROW_CAP:
            lines.append("…[行数过多已截断]…")
            break
        cells = [(c if isinstance(c, str) else str(c)).strip() for c in row]
        line = " | ".join(cells)
        lines.append(line)
        total += len(line) + 1
        if total >= lim:
            lines.append("…[体积过大已截断]…")
            break
    return "\n".join(lines)[:lim]


def _read_xlsx(path: Path, lim: int = _CHAR_CAP) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return ""
    parts: list[str] = []
    n = 0
    try:
        for sheet in wb.worksheets:
            parts.append(f"=== Sheet: {sheet.title} ===")
            for ri, row in enumerate(sheet.iter_rows(values_only=True)):
                if ri >= _ROW_CAP:
                    parts.append("…[行数过多已截断]…")
                    break
                cells = ["" if c is None else str(c).strip() for c in row]
                line = " | ".join(cells)
                parts.append(line)
                n += len(line) + 1
                if n >= lim:
                    parts.append("…[体积过大已截断]…")
                    return "\n".join(parts)[:lim]
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return "\n".join(parts)[:lim]


def extract_text_for_issuer(path: Path) -> str:
    e = path.suffix.lower()
    if e in (".txt", ".md", ".markdown"):
        return _read_plain(path)
    if e == ".pdf":
        return _read_pdf(path)
    if e == ".csv":
        return _read_csv(path)
    if e in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    if e == ".xls":
        return ""
    return ""
